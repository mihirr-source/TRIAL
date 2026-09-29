"""Tender document text extraction (PDF / DOCX).

Used by `POST /api/analyze/file`. Keeps the same philosophy as the rest of the
engine: no external services, clear errors, hard caps.

  * PDF: pdfplumber (text layer only). Scanned/image-only PDFs raise a clear
    error — OCR is deliberately out of scope for this release.
  * DOCX: python-docx, paragraphs + tables.
  * Caps: 10 MB per file, 50 PDF pages.
"""

from __future__ import annotations

import io

MAX_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_PAGES = 50

SUPPORTED_EXTENSIONS = (".pdf", ".docx")


class ExtractionError(ValueError):
    """Raised with a user-presentable message for any extraction failure."""


def extract_pdf(data: bytes) -> tuple[str, int]:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - guarded by requirements
        raise ExtractionError("PDF support unavailable — install the requirements.txt extras.") from exc
    try:
        pages_text: list[str] = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            total = len(pdf.pages)
            for page in pdf.pages[:MAX_PAGES]:
                pages_text.append(page.extract_text() or "")
        text = "\n".join(pages_text)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("Could not read the PDF file — it may be corrupt or encrypted.") from exc
    if not text.strip():
        raise ExtractionError(
            "No text layer found — the PDF looks scanned/image-only. "
            "Paste the specification text instead (OCR is not supported).")
    return text, total


def extract_docx(data: bytes) -> tuple[str, int | None]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - guarded by requirements
        raise ExtractionError("DOCX support unavailable — install the requirements.txt extras.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("Could not read the DOCX file — it may be corrupt or password-protected.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:  # BOQ schedules live in tables
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    text = "\n".join(parts)
    if not text.strip():
        raise ExtractionError("The DOCX contains no extractable text.")
    return text, None


def extract_text(data: bytes, filename: str) -> tuple[str, int | None]:
    """Extract text from an upload. Returns (text, page_count_or_None)."""
    if len(data) > MAX_BYTES:
        raise ExtractionError("File exceeds the 10 MB limit.")
    if not data:
        raise ExtractionError("The uploaded file is empty.")
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return extract_pdf(data)
    if name.endswith(".docx"):
        return extract_docx(data)
    raise ExtractionError("Unsupported file type — upload a PDF or DOCX file.")
